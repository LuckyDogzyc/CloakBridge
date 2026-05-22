from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from cloakbridge.main import app
from fixtures import write_minimal_docx
from test_docx_processor import read_docx_text


def test_api_sanitizes_multiple_files_with_shared_token_map(tmp_path: Path, monkeypatch):
    monkeypatch.setenv("CLOAKBRIDGE_DATA_DIR", str(tmp_path / "data"))
    txt = tmp_path / "input.txt"
    docx = tmp_path / "input.docx"
    xlsx = tmp_path / "input.xlsx"
    txt.write_text("西调工程服务器10.18.2.18", encoding="utf-8")
    write_minimal_docx(docx, "<w:p><w:r><w:t>复核西调工程</w:t></w:r></w:p>")
    workbook = Workbook()
    workbook.active["A1"] = "西调工程"
    workbook.save(xlsx)

    with txt.open("rb") as txt_file, docx.open("rb") as docx_file, xlsx.open("rb") as xlsx_file:
        response = TestClient(app).post(
            "/api/jobs/sanitize-files",
            files=[
                ("files", ("input.txt", txt_file, "text/plain")),
                ("files", ("input.docx", docx_file, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")),
                ("files", ("input.xlsx", xlsx_file, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")),
            ],
            data={
                "alias_groups": (
                    '[{"entity_type":"PROJECT","canonical":"西调工程",'
                    '"aliases":["西调工程","西调搬迁"]}]'
                )
            },
        )

    assert response.status_code == 200
    payload = response.json()
    assert payload["token_map"]["[[PRJ:001#001]]"] == "西调工程"
    assert len(payload["files"]) == 3

    txt_output = next(item for item in payload["files"] if item["filename"] == "input.txt")
    docx_output = next(item for item in payload["files"] if item["filename"] == "input.docx")
    xlsx_output = next(item for item in payload["files"] if item["filename"] == "input.xlsx")

    assert Path(txt_output["output_path"]).read_text(encoding="utf-8") == (
        "[[PRJ:001#001]]服务器[[IP:A.B.C.018]]"
    )
    assert "[[PRJ:001#001]]" in read_docx_text(Path(docx_output["output_path"]))
    assert load_workbook(Path(xlsx_output["output_path"])).active["A1"].value == "[[PRJ:001#001]]"


def test_api_validates_ai_response_tokens_before_restore():
    token_map = {
        "[[PRJ:001]]": "西调工程",
        "[[PRJ:001#001]]": "西调工程",
        "[[IP:A.B.C.018]]": "10.18.2.18",
    }

    response = TestClient(app).post(
        "/api/validate-response",
        json={
            "sanitized_text": "[[PRJ:001]]缺少[[IP:A.B.C.018]]，新增[[PRJ:999#001]]，坏标记[[PRJ001]]",
            "token_map": token_map,
        },
    )

    assert response.status_code == 200
    payload = response.json()
    assert payload["generic_tokens"] == ["[[PRJ:001]]"]
    assert payload["unknown_tokens"] == ["[[PRJ:999#001]]"]
    assert payload["malformed_tokens"] == ["[[PRJ001]]"]
    assert payload["restored_text"].startswith("西调工程缺少10.18.2.18")
