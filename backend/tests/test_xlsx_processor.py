from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from cloakbridge.documents.xlsx_processor import XlsxProcessor
from cloakbridge.domain.entities import EntityType, Finding
from cloakbridge.domain.tokens import TokenMap


def test_xlsx_processor_replaces_strings_and_preserves_style(tmp_path: Path):
    source = tmp_path / "input.xlsx"
    output = tmp_path / "sanitized.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "华东三期项目"
    sheet["A1"] = "上海明远科技有限公司"
    sheet["A1"].font = Font(bold=True, color="FF0000")
    sheet["A1"].fill = PatternFill(fill_type="solid", fgColor="FFFF00")
    workbook.save(source)

    findings = [
        Finding("华东三期项目", EntityType.PROJECT, 0, 6, "user", 1.0),
        Finding("上海明远科技有限公司", EntityType.COMPANY, 0, 10, "user", 1.0),
    ]
    token_map = TokenMap()

    XlsxProcessor().sanitize(source, output, findings, token_map)

    result = load_workbook(output)
    result_sheet = result["<PROJECT_001>"]
    assert result_sheet["A1"].value == "<COMPANY_001>"
    assert result_sheet["A1"].font.bold is True
    assert result_sheet["A1"].font.color.rgb == "00FF0000"
    assert result_sheet["A1"].fill.fgColor.rgb == "00FFFF00"


def test_xlsx_processor_reuses_shared_token_map_across_workbooks(tmp_path: Path):
    first = tmp_path / "first.xlsx"
    second = tmp_path / "second.xlsx"
    first_output = tmp_path / "first.sanitized.xlsx"
    second_output = tmp_path / "second.sanitized.xlsx"
    token_map = TokenMap()

    first_workbook = Workbook()
    first_workbook.active["A1"] = "华东三期项目"
    first_workbook.save(first)

    second_workbook = Workbook()
    second_workbook.active["A1"] = "复核华东三期项目"
    second_workbook.save(second)

    XlsxProcessor().sanitize(
        first,
        first_output,
        [Finding("华东三期项目", EntityType.PROJECT, 0, 6, "user", 1.0)],
        token_map,
    )
    XlsxProcessor().sanitize(
        second,
        second_output,
        [Finding("华东三期项目", EntityType.PROJECT, 2, 8, "user", 1.0)],
        token_map,
    )

    assert load_workbook(first_output).active["A1"].value == "<PROJECT_001>"
    assert load_workbook(second_output).active["A1"].value == "复核<PROJECT_001>"
