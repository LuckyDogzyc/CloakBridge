from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from cloakbridge.domain.entities import Finding
from cloakbridge.domain.tokens import TokenMap


class XlsxProcessor:
    def sanitize(self, source: Path, output: Path, findings: list[Finding], token_map: TokenMap) -> None:
        replacements = {finding.text: token_map.token_for(finding) for finding in findings}
        workbook = load_workbook(source)
        for sheet in workbook.worksheets:
            sheet.title = self._replace(sheet.title, replacements)
            for row in sheet.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str):
                        cell.value = self._replace(cell.value, replacements)
                    if cell.comment and cell.comment.text:
                        cell.comment.text = self._replace(cell.comment.text, replacements)
                    if cell.hyperlink and cell.hyperlink.display:
                        cell.hyperlink.display = self._replace(cell.hyperlink.display, replacements)
        workbook.save(output)

    def _replace(self, text: str, replacements: dict[str, str]) -> str:
        result = text
        for original, token in sorted(replacements.items(), key=lambda item: -len(item[0])):
            result = result.replace(original, token)
        return result
