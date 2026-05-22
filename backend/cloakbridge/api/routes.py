from __future__ import annotations

import json
from dataclasses import asdict
import os
from pathlib import Path
import re
from uuid import uuid4
from zipfile import ZipFile

from fastapi import APIRouter, File, Form, HTTPException, UploadFile
from lxml import etree
from openpyxl import load_workbook
from pydantic import BaseModel

from cloakbridge.detection.dictionary import DictionaryEntry, DictionaryMatcher
from cloakbridge.detection.local_ai import HeuristicChineseLocalAI
from cloakbridge.detection.pipeline import DetectionPipeline
from cloakbridge.detection.regex_detector import RegexDetector
from cloakbridge.documents.docx_processor import DocxProcessor
from cloakbridge.documents.txt_processor import TxtProcessor
from cloakbridge.documents.xlsx_processor import XlsxProcessor
from cloakbridge.domain.entities import EntityType, Finding
from cloakbridge.domain.tokens import TokenMap
from cloakbridge.gateway.token_prompt import build_token_handling_prompt
from cloakbridge.gateway.providers import ProviderConfig
from cloakbridge.storage.sqlite_store import SQLiteStore
from cloakbridge.storage.vault import MappingVault

router = APIRouter(prefix="/api")

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
TOKEN_PATTERN = re.compile(r"\[\[[^\[\]]+\]\]")
VALID_TOKEN_PATTERN = re.compile(
    r"^\[\[(?:"
    r"(?:[A-Z_]+:\d{3}(?:#\d{3})?)|"
    r"(?:IP:[A-Z.]+\.\d{3}(?:/\d{1,2})?)|"
    r"(?:IP_PREFIX:[A-Z.]+\.\*)|"
    r"(?:IP_RANGE:[A-Z.]+\.\d{3}-\d{3})"
    r")\]\]$"
)
GENERIC_GROUP_PATTERN = re.compile(r"^\[\[[A-Z_]+:\d{3}\]\]$")


class DictionaryEntryIn(BaseModel):
    text: str
    entity_type: EntityType
    scope: str = "project"


class AnalyzeTextRequest(BaseModel):
    text: str
    dictionary: list[DictionaryEntryIn] = []


class FindingIn(BaseModel):
    text: str
    entity_type: EntityType
    start: int
    end: int
    source: str
    confidence: float


class AliasGroupIn(BaseModel):
    entity_type: EntityType
    canonical: str
    aliases: list[str]


class AliasGroupCreateRequest(AliasGroupIn):
    scope: str = "project"


class SanitizeTextRequest(BaseModel):
    text: str
    findings: list[FindingIn]
    alias_groups: list[AliasGroupIn] = []


class RestoreTextRequest(BaseModel):
    sanitized_text: str
    token_map: dict[str, str]


class ValidateResponseRequest(RestoreTextRequest):
    pass


class ModelConfigCreateRequest(BaseModel):
    name: str
    provider: str
    model: str
    base_url: str = ""
    api_key: str = ""


@router.post("/analyze-text")
def analyze_text(request: AnalyzeTextRequest) -> dict[str, object]:
    entries = [
        DictionaryEntry(item.text, item.entity_type, item.scope, "confirmed")
        for item in request.dictionary
    ]
    pipeline = DetectionPipeline(RegexDetector(), DictionaryMatcher(entries), HeuristicChineseLocalAI())
    findings = pipeline.detect(request.text)
    return {"findings": [asdict(finding) for finding in findings]}


@router.post("/alias-groups")
def create_alias_group(request: AliasGroupCreateRequest) -> dict[str, object]:
    store = _local_store()
    return store.create_alias_group(
        entity_type=request.entity_type.value,
        canonical=request.canonical,
        aliases=request.aliases,
        scope=request.scope,
    )


@router.get("/alias-groups")
def list_alias_groups(scope: str | None = None) -> dict[str, object]:
    store = _local_store()
    return {"alias_groups": store.list_alias_groups(scope)}


@router.post("/model-configs")
def create_model_config(request: ModelConfigCreateRequest) -> dict[str, object]:
    try:
        ProviderConfig(provider=request.provider, model=request.model, base_url=request.base_url or None)
    except ValueError as error:
        raise HTTPException(status_code=400, detail=str(error)) from error

    store = _local_store()
    config = store.create_model_config(
        name=request.name.strip(),
        provider=request.provider,
        model=request.model.strip(),
        base_url=request.base_url.strip(),
    )
    if request.api_key.strip():
        secret_ref = f"model-config-{config['id']}"
        _local_vault().save_secret(secret_ref, request.api_key.strip())
        config = store.update_model_config_secret(
            int(config["id"]),
            secret_ref,
            _mask_api_key(request.api_key.strip()),
        )
    return _public_model_config(config)


@router.get("/model-configs")
def list_model_configs() -> dict[str, object]:
    store = _local_store()
    return {"model_configs": [_public_model_config(config) for config in store.list_model_configs()]}


@router.post("/model-configs/{config_id}/test")
def test_model_config(config_id: int) -> dict[str, object]:
    store = _local_store()
    try:
        config = store.get_model_config(config_id)
    except KeyError as error:
        raise HTTPException(status_code=404, detail="Model config not found") from error

    if not config["model"]:
        raise HTTPException(status_code=400, detail="Model id is required")
    if not config["base_url"]:
        raise HTTPException(status_code=400, detail="Base URL is required")
    if not config["secret_ref"]:
        raise HTTPException(status_code=400, detail="API key is required")

    _local_vault().load_secret(str(config["secret_ref"]))
    return {"ok": True, "message": "配置可用", "provider": config["provider"], "model": config["model"]}


@router.post("/sanitize-text")
def sanitize_text(request: SanitizeTextRequest) -> dict[str, object]:
    token_map = TokenMap(replacement_style="structured")
    for group in request.alias_groups:
        token_map.register_alias_group(group.entity_type, group.canonical, group.aliases)
    findings = [
        Finding(item.text, item.entity_type, item.start, item.end, item.source, item.confidence)
        for item in request.findings
    ]
    sanitized_text = TxtProcessor().replace_text(request.text, findings, token_map)
    return {
        "sanitized_text": sanitized_text,
        "token_map": token_map.token_to_original,
        "token_prompt": build_token_handling_prompt(token_map.token_to_original),
    }


@router.post("/jobs/sanitize-files")
async def sanitize_files(
    files: list[UploadFile] = File(...),
    alias_groups: str = Form("[]"),
) -> dict[str, object]:
    groups = _parse_alias_groups(alias_groups)
    token_map = TokenMap(replacement_style="structured")
    for group in groups:
        token_map.register_alias_group(group.entity_type, group.canonical, group.aliases)

    job_id = uuid4().hex
    data_dir = _data_dir()
    upload_dir = data_dir / "uploads" / job_id
    output_dir = data_dir / "outputs" / job_id
    upload_dir.mkdir(parents=True, exist_ok=True)
    output_dir.mkdir(parents=True, exist_ok=True)

    results = []
    for upload in files:
        filename = Path(upload.filename or "upload.bin").name
        source = upload_dir / filename
        source.write_bytes(await upload.read())
        text = _extract_file_text(source)
        findings = _pipeline_for_alias_groups(groups).detect(text)
        output = output_dir / filename
        _sanitize_file(source, output, findings, token_map)
        results.append(
            {
                "filename": filename,
                "output_path": str(output),
                "finding_count": len(findings),
            }
        )

    return {
        "job_id": job_id,
        "files": results,
        "token_map": token_map.token_to_original,
        "token_prompt": build_token_handling_prompt(token_map.token_to_original),
    }


@router.post("/restore-text")
def restore_text(request: RestoreTextRequest) -> dict[str, str]:
    restored = _restore_once(request.sanitized_text, request.token_map)
    return {"restored_text": restored}


@router.post("/validate-response")
def validate_response(request: ValidateResponseRequest) -> dict[str, object]:
    tokens = TOKEN_PATTERN.findall(request.sanitized_text)
    unknown_tokens = sorted(
        {token for token in tokens if VALID_TOKEN_PATTERN.fullmatch(token) and token not in request.token_map}
    )
    malformed_tokens = sorted({token for token in tokens if not VALID_TOKEN_PATTERN.fullmatch(token)})
    generic_tokens = sorted(
        {token for token in tokens if GENERIC_GROUP_PATTERN.fullmatch(token) and token in request.token_map}
    )
    return {
        "unknown_tokens": unknown_tokens,
        "malformed_tokens": malformed_tokens,
        "generic_tokens": generic_tokens,
        "restored_text": _restore_once(request.sanitized_text, request.token_map),
    }


def _restore_once(text: str, token_map: dict[str, str]) -> str:
    if not token_map:
        return text
    pattern = re.compile("|".join(re.escape(token) for token in sorted(token_map, key=len, reverse=True)))
    return pattern.sub(lambda match: token_map[match.group(0)], text)


def _local_store() -> SQLiteStore:
    store = SQLiteStore(_data_dir() / "cloakbridge.sqlite")
    store.initialize()
    return store


def _local_vault() -> MappingVault:
    return MappingVault(_data_dir() / "vault")


def _data_dir() -> Path:
    return Path(os.environ.get("CLOAKBRIDGE_DATA_DIR", ".local-data")).resolve()


def _mask_api_key(api_key: str) -> str:
    if len(api_key) <= 8:
        return f"{api_key[:2]}****"
    return f"{api_key[:3]}****{api_key[-4:]}"


def _public_model_config(config: dict[str, object]) -> dict[str, object]:
    return {
        "id": config["id"],
        "name": config["name"],
        "provider": config["provider"],
        "model": config["model"],
        "base_url": config["base_url"],
        "masked_api_key": config["masked_api_key"],
        "enabled": config["enabled"],
    }


def _parse_alias_groups(payload: str) -> list[AliasGroupIn]:
    raw_groups = json.loads(payload)
    return [AliasGroupIn(**item) for item in raw_groups]


def _pipeline_for_alias_groups(alias_groups: list[AliasGroupIn]) -> DetectionPipeline:
    dictionary_entries = [
        DictionaryEntry(alias, group.entity_type, "job", "confirmed")
        for group in alias_groups
        for alias in group.aliases
    ]
    return DetectionPipeline(
        RegexDetector(),
        DictionaryMatcher(dictionary_entries),
        HeuristicChineseLocalAI(),
    )


def _extract_file_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".txt":
        return path.read_text(encoding="utf-8")
    if suffix == ".docx":
        with ZipFile(path) as archive:
            xml = archive.read("word/document.xml")
        root = etree.fromstring(xml)
        return "".join(root.xpath("//w:t/text()", namespaces={"w": W_NS}))
    if suffix == ".xlsx":
        workbook = load_workbook(path)
        parts: list[str] = []
        for sheet in workbook.worksheets:
            parts.append(sheet.title)
            for row in sheet.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str):
                        parts.append(cell.value)
        return "\n".join(parts)
    raise ValueError(f"Unsupported file type: {path.suffix}")


def _sanitize_file(source: Path, output: Path, findings: list[Finding], token_map: TokenMap) -> None:
    suffix = source.suffix.lower()
    if suffix == ".txt":
        sanitized = TxtProcessor().sanitize(source, findings, token_map)
        output.write_text(sanitized.text, encoding=sanitized.encoding, newline="")
        return
    if suffix == ".docx":
        DocxProcessor().sanitize(source, output, findings, token_map)
        return
    if suffix == ".xlsx":
        XlsxProcessor().sanitize(source, output, findings, token_map)
        return
    raise ValueError(f"Unsupported file type: {source.suffix}")
